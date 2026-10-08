from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from table_gen.config import Config
from table_gen.dates import format_header_date, is_date_like


def generate_attendance_markdown(wb: Workbook, sheet_name: str, config: Config) -> str:
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet_name}' not found in the Excel file.")

    sheet = wb[sheet_name]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index_col = config.student_columns.index

    lecture_cols = [
        i
        for i, val in enumerate(header_row)
        if i > index_col and is_date_like(val)
    ]
    if not lecture_cols:
        raise ValueError(f"No date columns found in sheet '{sheet_name}'.")

    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) <= index_col:
            continue
        nr_indeksu = row[index_col]
        if nr_indeksu is None or str(nr_indeksu).strip() == "":
            continue
        # Skip accidental non-numeric junk (e.g. pasted HTML)
        if not str(nr_indeksu).strip().isdigit():
            continue
        rows.append(row)

    rows.sort(key=lambda r: str(r[index_col]).strip())

    header_cells = ["Nr indeksu"] + [format_header_date(header_row[i]) for i in lecture_cols]

    md = (
        "<style>\n"
        "\ttable, th, td {\n"
        "\t\tborder: 1px solid black;\n"
        "\t\tborder-collapse: collapse;\n"
        "\t}\n"
        "</style>\n\n"
    )
    md += "| " + " | ".join(header_cells) + " |\n"
    md += "| " + " | ".join(["---"] * len(header_cells)) + " |\n"

    for row in rows:
        cells = [str(row[index_col]).strip()]
        for col in lecture_cols:
            value = row[col] if len(row) > col else None
            cells.append("+" if value else "")
        md += "| " + " | ".join(cells) + " |\n"

    return md


def write_attendance_tables(config: Config) -> list[str]:
    if not config.input_file.is_file():
        raise FileNotFoundError(f"Input Excel not found: {config.input_file}")

    config.output_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(filename=config.input_file)
    written: list[str] = []

    for seminar in config.enabled_seminars:
        if seminar.name not in wb.sheetnames:
            print(f"Skipping '{seminar.name}': sheet not found in workbook.")
            continue
        try:
            markdown = generate_attendance_markdown(wb, seminar.name, config)
        except ValueError as exc:
            print(f"Skipping '{seminar.name}': {exc}")
            continue

        safe_name = seminar.name.replace(" ", "_")
        output_file = config.output_dir / f"{safe_name}_attendance.md"
        output_file.write_text(markdown, encoding="utf-8")
        written.append(str(output_file))
        print(f"Saved: {output_file}")

    return written
