from __future__ import annotations

from datetime import datetime, timedelta

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from table_gen.config import Config
from table_gen.dates import parse_date


def load_lookup(sheet: Worksheet, email_col: int, value_col: int) -> dict[str, object]:
    """Return dict: email -> value (date string or title string)."""
    lookup: dict[str, object] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) <= max(email_col, value_col):
            continue
        email_raw = row[email_col]
        if email_raw is None:
            continue
        email = str(email_raw).strip()
        if not email or "@" not in email:
            continue
        lookup[email] = row[value_col]
    return lookup


def generate_weekly_schedule(wb: Workbook, sheet_name: str, config: Config) -> str:
    dates_sheet_name = sheet_name + config.dates_suffix
    titles_sheet_name = sheet_name + config.titles_suffix

    missing = [n for n in (sheet_name, dates_sheet_name, titles_sheet_name) if n not in wb.sheetnames]
    if missing:
        raise KeyError(f"Missing sheet(s) for '{sheet_name}': {', '.join(missing)}")

    student_sheet = wb[sheet_name]
    date_lookup = load_lookup(
        wb[dates_sheet_name],
        config.lookup_columns.email,
        config.lookup_columns.dates_value,
    )
    title_lookup = load_lookup(
        wb[titles_sheet_name],
        config.lookup_columns.email,
        config.lookup_columns.titles_value,
    )

    cols = config.student_columns
    lectures = []

    for row in student_sheet.iter_rows(min_row=2, values_only=True):
        if len(row) <= cols.email:
            continue
        email_raw = row[cols.email]
        if email_raw is None:
            continue
        email = str(email_raw).strip()
        if not email or "@" not in email:
            continue

        name = row[cols.name] if len(row) > cols.name else ""
        surname = row[cols.surname] if len(row) > cols.surname else ""
        english_raw = row[cols.english] if len(row) > cols.english else ""
        english = str(english_raw or "").strip().lower() == "yes"

        lecture_date = parse_date(date_lookup.get(email))

        lecture_title = title_lookup.get(email, "")
        if lecture_title is None:
            lecture_title = ""
        lecture_title = str(lecture_title).strip()

        if lecture_date and lecture_title:
            lectures.append(
                {
                    "date": lecture_date,
                    "title": lecture_title,
                    "name": f"{name} {surname}".strip(),
                    "english": english,
                }
            )

    lectures.sort(key=lambda item: item["date"])

    start = datetime.strptime(config.schedule.start_date, "%Y-%m-%d").date()
    end = datetime.strptime(config.schedule.end_date, "%Y-%m-%d").date()
    step = timedelta(days=config.schedule.week_interval_days)

    md = ""
    current_date = start
    week_number = 1

    while current_date <= end:
        display_date = current_date.strftime("%d.%m.%Y")
        week_lectures = [lec for lec in lectures if lec["date"] == current_date]

        md += f"## {display_date}\n\n"
        if not week_lectures:
            md += "```\nWykład nie odbędzie się\n```\n\n"
        else:
            for lec in week_lectures:
                if lec["english"]:
                    line = (
                        f"Lecture {week_number}: **{lec['title']}** – "
                        f"*{lec['name']}* (open lecture) "
                        f"*Lecture will be held in English*"
                    )
                else:
                    line = (
                        f"Wykład {week_number}: **{lec['title']}** – "
                        f"*{lec['name']}* (wykład otwarty)"
                    )
                md += f"```\n{line}\n```\n\n"
                week_number += 1

        current_date += step

    return md


def write_weekly_schedules(config: Config) -> list[str]:
    if not config.input_file.is_file():
        raise FileNotFoundError(f"Input Excel not found: {config.input_file}")

    config.output_dir.mkdir(parents=True, exist_ok=True)
    # data_only=True reads cached formula results (e.g. date column F).
    # Re-save the workbook in Excel after changing formulas so values are cached.
    wb = load_workbook(filename=config.input_file, data_only=True)
    written: list[str] = []

    for seminar in config.enabled_seminars:
        try:
            md = generate_weekly_schedule(wb, seminar.name, config)
        except KeyError as exc:
            print(f"Skipping '{seminar.name}': {exc}")
            continue

        safe_name = seminar.name.replace(" ", "_")
        output_file = config.output_dir / f"{safe_name}_weekly_schedule.md"
        output_file.write_text(md, encoding="utf-8")
        written.append(str(output_file))
        print(f"Saved schedule: {output_file}")

    return written
